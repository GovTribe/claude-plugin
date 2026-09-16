#!/usr/bin/env python3
"""Convert a GovTribe deep-dive Markdown dossier into a polished DOCX.

The converter intentionally supports the Markdown subset used by the deep-dive
output contract: ATX headings, paragraphs, bullet/numbered lists, blockquotes,
code fences, and pipe tables. Markdown remains the source of truth.
"""

from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass
from datetime import date
from pathlib import Path


def load_python_docx() -> str | None:
    """Lazy-load the optional DOCX dependency for the conversion operation."""
    global Document
    global WD_ORIENT, WD_STYLE_TYPE, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    global WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
    global OxmlElement, qn, Inches, Pt, RGBColor

    try:
        from docx import Document as DocxDocument
        from docx.enum.section import WD_ORIENT as DocxOrientation
        from docx.enum.style import WD_STYLE_TYPE as DocxStyleType
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT as DocxCellVerticalAlignment
        from docx.enum.table import WD_TABLE_ALIGNMENT as DocxTableAlignment
        from docx.enum.text import WD_ALIGN_PARAGRAPH as DocxParagraphAlignment
        from docx.enum.text import WD_BREAK as DocxBreak
        from docx.enum.text import WD_TAB_ALIGNMENT as DocxTabAlignment
        from docx.oxml import OxmlElement as DocxOxmlElement
        from docx.oxml.ns import qn as docx_qn
        from docx.shared import Inches as DocxInches
        from docx.shared import Pt as DocxPoints
        from docx.shared import RGBColor as DocxRgbColor
    except ImportError:
        return (
            "optional dependency python-docx is unavailable; the validated Markdown remains usable. "
            "Use the host's document-generation capability or run this converter in an environment "
            "where python-docx is already provided."
        )

    Document = DocxDocument
    WD_ORIENT = DocxOrientation
    WD_STYLE_TYPE = DocxStyleType
    WD_CELL_VERTICAL_ALIGNMENT = DocxCellVerticalAlignment
    WD_TABLE_ALIGNMENT = DocxTableAlignment
    WD_ALIGN_PARAGRAPH = DocxParagraphAlignment
    WD_BREAK = DocxBreak
    WD_TAB_ALIGNMENT = DocxTabAlignment
    OxmlElement = DocxOxmlElement
    qn = docx_qn
    Inches = DocxInches
    Pt = DocxPoints
    RGBColor = DocxRgbColor

    return None


@dataclass(frozen=True)
class StylePreset:
    body_font: str
    heading_font: str
    body_size: float
    title_size: float
    heading_color: str
    accent_color: str
    table_header_fill: str
    callout_fill: str


STYLE_PRESETS: dict[str, StylePreset] = {
    "govtribe": StylePreset(
        body_font="Aptos",
        heading_font="Aptos Display",
        body_size=10.0,
        title_size=24.0,
        heading_color="17365D",
        accent_color="2F5597",
        table_header_fill="D9EAF7",
        callout_fill="EAF2F8",
    ),
    "plain": StylePreset(
        body_font="Arial",
        heading_font="Arial",
        body_size=10.0,
        title_size=22.0,
        heading_color="222222",
        accent_color="444444",
        table_header_fill="E7E6E6",
        callout_fill="F2F2F2",
    ),
}

PIPE_SEPARATOR_RE = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(?:\|\s*:?-{3,}:?\s*)+\|?\s*$")
HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
BULLET_RE = re.compile(r"^\s*[-*+]\s+(.+)$")
NUMBERED_RE = re.compile(r"^\s*\d+[.)]\s+(.+)$")
INLINE_TOKEN_RE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|__[^_]+__|\*[^*]+\*|_[^_]+_|\[[^\]]+\]\([^)]+\))")


def hex_color(value: str) -> RGBColor:
    value = value.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def set_cell_fill(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(end)


def add_inline_runs(paragraph, text: str, *, font_name: str | None = None, font_size: float | None = None) -> None:
    cursor = 0
    for match in INLINE_TOKEN_RE.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            if font_name:
                run.font.name = font_name
            if font_size:
                run.font.size = Pt(font_size)

        token = match.group(0)
        display = token
        bold = italic = False
        code = False

        if token.startswith("**") or token.startswith("__"):
            display = token[2:-2]
            bold = True
        elif token.startswith("`"):
            display = token[1:-1]
            code = True
        elif token.startswith("["):
            label, url = token[1:].split("](", 1)
            display = f"{label} ({url[:-1]})"
        else:
            display = token[1:-1]
            italic = True

        run = paragraph.add_run(display)
        run.bold = bold
        run.italic = italic
        if code:
            run.font.name = "Consolas"
            run.font.size = Pt((font_size or 9.0) - 0.5)
        else:
            if font_name:
                run.font.name = font_name
            if font_size:
                run.font.size = Pt(font_size)
        cursor = match.end()

    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        if font_name:
            run.font.name = font_name
        if font_size:
            run.font.size = Pt(font_size)


def parse_pipe_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def configure_document(document: Document, preset: StylePreset, *, landscape: bool) -> None:
    for section in document.sections:
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)
        if landscape:
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width, section.page_height = section.page_height, section.page_width

    normal = document.styles["Normal"]
    normal.font.name = preset.body_font
    normal.font.size = Pt(preset.body_size)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.05

    for level in range(1, 7):
        style = document.styles[f"Heading {level}"]
        style.font.name = preset.heading_font
        style.font.color.rgb = hex_color(preset.heading_color)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10 if level <= 2 else 7)
        style.paragraph_format.space_after = Pt(4)

    document.styles["Title"].font.name = preset.heading_font
    document.styles["Title"].font.size = Pt(preset.title_size)
    document.styles["Title"].font.color.rgb = hex_color(preset.heading_color)

    if "GovTribe Callout" not in document.styles:
        callout = document.styles.add_style("GovTribe Callout", WD_STYLE_TYPE.PARAGRAPH)
        callout.base_style = document.styles["Normal"]
        callout.font.name = preset.body_font
        callout.font.size = Pt(preset.body_size)
        callout.font.color.rgb = hex_color(preset.heading_color)
        callout.paragraph_format.left_indent = Inches(0.2)
        callout.paragraph_format.right_indent = Inches(0.1)
        callout.paragraph_format.space_before = Pt(4)
        callout.paragraph_format.space_after = Pt(6)

    if "GovTribe Evidence" not in document.styles:
        evidence = document.styles.add_style("GovTribe Evidence", WD_STYLE_TYPE.PARAGRAPH)
        evidence.base_style = document.styles["Normal"]
        evidence.font.name = preset.body_font
        evidence.font.size = Pt(8.5)
        evidence.font.color.rgb = RGBColor(80, 80, 80)


def shade_paragraph(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_footer(document: Document, *, source_note: str, generated_date: str, page_numbers: bool) -> None:
    for section in document.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        paragraph = footer.paragraphs[0]
        paragraph.clear()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(3.15), WD_TAB_ALIGNMENT.CENTER)
        paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.25), WD_TAB_ALIGNMENT.RIGHT)

        left = paragraph.add_run(source_note)
        left.font.size = Pt(8)
        left.font.color.rgb = RGBColor(90, 90, 90)

        middle = paragraph.add_run(f"\tGenerated {generated_date}")
        middle.font.size = Pt(8)
        middle.font.color.rgb = RGBColor(90, 90, 90)

        if page_numbers:
            right = paragraph.add_run("\tPage ")
            right.font.size = Pt(8)
            right.font.color.rgb = RGBColor(90, 90, 90)
            add_page_number(paragraph)


def set_document_metadata(document: Document, *, title: str, generated_date: str) -> None:
    document.core_properties.title = title
    document.core_properties.subject = "GovTribe Deep Dive"
    document.core_properties.author = "GovTribe"
    document.core_properties.comments = f"Generated {generated_date} from a validated Markdown source."


def add_markdown_table(document: Document, rows: list[list[str]], preset: StylePreset, *, evidence_ledger: bool) -> None:
    if not rows:
        return
    columns = max(len(row) for row in rows)
    normalized = [row + [""] * (columns - len(row)) for row in rows]
    table = document.add_table(rows=len(normalized), cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False

    usable_width = 9.4 if document.sections[0].orientation == WD_ORIENT.LANDSCAPE else 6.05
    column_width = Inches(usable_width / max(columns, 1))

    for row_index, values in enumerate(normalized):
        row = table.rows[row_index]
        if row_index == 0:
            set_repeat_table_header(row)
        for col_index, value in enumerate(values):
            cell = row.cells[col_index]
            cell.width = column_width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline_runs(
                paragraph,
                value,
                font_name=preset.body_font,
                font_size=8.0 if evidence_ledger else 8.5,
            )
            if row_index == 0:
                set_cell_fill(cell, preset.table_header_fill)
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = hex_color(preset.heading_color)

    document.add_paragraph()


def add_code_block(document: Document, lines: list[str]) -> None:
    paragraph = document.add_paragraph()
    paragraph.style = document.styles["GovTribe Evidence"]
    paragraph.paragraph_format.left_indent = Inches(0.2)
    paragraph.paragraph_format.right_indent = Inches(0.2)
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(8)
        if index < len(lines) - 1:
            run.add_break()


def convert_markdown(document: Document, text: str, preset: StylePreset, *, title_override: str | None) -> str:
    lines = text.splitlines()
    title = title_override or "GovTribe Deep Dive"
    index = 0
    current_heading = ""
    code_fence = False
    code_lines: list[str] = []

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            if code_fence:
                add_code_block(document, code_lines)
                code_lines = []
                code_fence = False
            else:
                code_fence = True
            index += 1
            continue

        if code_fence:
            code_lines.append(line)
            index += 1
            continue

        if not stripped:
            index += 1
            continue

        heading = HEADING_RE.match(line)
        if heading:
            level = len(heading.group(1))
            heading_text = heading.group(2).strip()
            current_heading = heading_text.lower()
            if level == 1 and title == "GovTribe Deep Dive":
                title = re.sub(r"[*_`]", "", heading_text)
                paragraph = document.add_paragraph(style="Title")
                add_inline_runs(paragraph, heading_text, font_name=preset.heading_font)
            elif level == 1:
                paragraph = document.add_paragraph(style="Title")
                add_inline_runs(paragraph, title_override or heading_text, font_name=preset.heading_font)
            else:
                paragraph = document.add_heading(level=min(level - 1, 6))
                add_inline_runs(paragraph, heading_text, font_name=preset.heading_font)
            index += 1
            continue

        if index + 1 < len(lines) and "|" in line and PIPE_SEPARATOR_RE.match(lines[index + 1]):
            table_rows = [parse_pipe_row(line)]
            index += 2
            while index < len(lines) and "|" in lines[index] and lines[index].strip():
                table_rows.append(parse_pipe_row(lines[index]))
                index += 1
            add_markdown_table(
                document,
                table_rows,
                preset,
                evidence_ledger="evidence ledger" in current_heading,
            )
            continue

        bullet = BULLET_RE.match(line)
        if bullet:
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline_runs(paragraph, bullet.group(1), font_name=preset.body_font, font_size=preset.body_size)
            index += 1
            continue

        numbered = NUMBERED_RE.match(line)
        if numbered:
            paragraph = document.add_paragraph(style="List Number")
            add_inline_runs(paragraph, numbered.group(1), font_name=preset.body_font, font_size=preset.body_size)
            index += 1
            continue

        if stripped.startswith(">"):
            paragraph = document.add_paragraph(style="GovTribe Callout")
            shade_paragraph(paragraph, preset.callout_fill)
            add_inline_runs(
                paragraph,
                stripped.lstrip("> "),
                font_name=preset.body_font,
                font_size=preset.body_size,
            )
            index += 1
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index]
            candidate_stripped = candidate.strip()
            if not candidate_stripped:
                break
            if (
                HEADING_RE.match(candidate)
                or BULLET_RE.match(candidate)
                or NUMBERED_RE.match(candidate)
                or candidate_stripped.startswith(">")
                or candidate_stripped.startswith("```")
                or (index + 1 < len(lines) and "|" in candidate and PIPE_SEPARATOR_RE.match(lines[index + 1]))
            ):
                break
            paragraph_lines.append(candidate_stripped)
            index += 1

        paragraph = document.add_paragraph()
        add_inline_runs(
            paragraph,
            " ".join(paragraph_lines),
            font_name=preset.body_font,
            font_size=preset.body_size,
        )

    if code_fence and code_lines:
        add_code_block(document, code_lines)

    return title


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Convert GovTribe deep-dive Markdown to DOCX.")
    parser.add_argument("markdown", type=Path, help="Input Markdown dossier.")
    parser.add_argument("docx", type=Path, help="Output DOCX path.")
    parser.add_argument("--style", choices=sorted(STYLE_PRESETS), default="govtribe")
    parser.add_argument("--title", help="Optional document title override.")
    parser.add_argument(
        "--source-note",
        default="GovTribe Deep Dive - source details in evidence ledger",
        help="Footer source note.",
    )
    parser.add_argument(
        "--generated-date",
        default=date.today().isoformat(),
        help="Generated date displayed in the footer. Default: today.",
    )
    parser.add_argument("--landscape", action="store_true", help="Use landscape orientation.")
    parser.add_argument("--no-page-numbers", action="store_true", help="Omit footer page numbers.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if not args.markdown.is_file():
        print(f"error: Markdown file not found: {args.markdown}", file=sys.stderr)
        return 2

    dependency_error = load_python_docx()
    if dependency_error:
        print(f"error: DOCX output unavailable: {dependency_error}", file=sys.stderr)
        return 3

    try:
        text = args.markdown.read_text(encoding="utf-8")
        document = Document()
        preset = STYLE_PRESETS[args.style]
        configure_document(document, preset, landscape=args.landscape)
        title = convert_markdown(document, text, preset, title_override=args.title)
        set_document_metadata(document, title=title, generated_date=args.generated_date)
        add_footer(
            document,
            source_note=args.source_note,
            generated_date=args.generated_date,
            page_numbers=not args.no_page_numbers,
        )
        args.docx.parent.mkdir(parents=True, exist_ok=True)
        document.save(args.docx)
    except Exception as exc:  # noqa: BLE001 - CLI should return a useful failure.
        print(f"error: failed to create DOCX: {exc}", file=sys.stderr)
        return 1

    print(f"created DOCX: {args.docx}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
