#!/usr/bin/env python3
"""Validate a GovTribe proposal control workbook smoke artifact."""
from __future__ import annotations

import json
import sys
import argparse
from pathlib import Path

EXPECTED_SHEETS = [
    "Start Here",
    "Opportunity Setup",
    "Dashboard",
    "Requirement Matrix",
    "Evaluation Crosswalk",
    "Submission Checklist",
    "Pricing & Deliverables",
    "Questions-Risks",
    "Amendment Log",
    "Sources",
    "Lists",
]

HEADER_ROWS = {
    "Requirement Matrix": 7,
    "Evaluation Crosswalk": 7,
    "Submission Checklist": 6,
    "Pricing & Deliverables": 6,
    "Questions-Risks": 6,
    "Amendment Log": 6,
    "Sources": 5,
}

REQUIRED_HEADERS = {
    "Requirement Matrix": {
        "Req ID",
        "Source Area",
        "Requirement Category",
        "Requirement Level",
        "Source Citation",
        "Verbatim Requirement",
        "Proposal Action / Interpretation",
        "Proposal Volume",
        "Proposal Outline Ref",
        "Owner",
        "Compliance Posture",
        "Draft Status",
        "Evidence / Artifact Needed",
        "Control Flag",
    },
    "Evaluation Crosswalk": {"Eval ID", "Source Citation", "Factor / Subfactor", "Related Req IDs"},
    "Submission Checklist": {"Check ID", "Category", "Control Item", "Status", "Control Flag"},
    "Pricing & Deliverables": {"Item ID", "Item Type", "Source Citation", "Description", "Related Requirement IDs"},
    "Questions-Risks": {"Item ID", "Type", "Priority", "Issue / Decision", "Status"},
    "Amendment Log": {"Amendment ID", "Date", "Change Type", "Summary of Change / Clarification", "Impacted Req IDs"},
    "Sources": {"Source Type", "Short Name", "What it supports", "Key Takeaway / Notes", "Source URL"},
}

CORE_REQUIREMENT_COLUMNS = {
    "Req ID",
    "Source Area",
    "Requirement Level",
    "Source Citation",
    "Verbatim Requirement",
    "Proposal Action / Interpretation",
    "Proposal Volume",
    "Proposal Outline Ref",
    "Owner",
    "Compliance Posture",
    "Draft Status",
    "Evidence / Artifact Needed",
}


def load_openpyxl_workbook(path: Path, data_only: bool = False):
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as exc:
        print(
            "DEGRADED: openpyxl is unavailable. Validate the workbook manually against "
            "references/workbook-schema.md and references/quality-checks.md, and "
            "deliver Markdown or CSV control views when workbook validation cannot run.",
            file=sys.stderr,
        )
        raise SystemExit(2) from exc

    return load_workbook(path, data_only=data_only)


def populated(value: object) -> bool:
    return value is not None and str(value).strip() != ""


def row_values(ws, row_idx: int) -> list[object]:
    return [ws.cell(row=row_idx, column=col_idx).value for col_idx in range(1, ws.max_column + 1)]


def header_map(ws) -> dict[str, int]:
    row_idx = HEADER_ROWS[ws.title]
    return {
        str(value): idx
        for idx, value in enumerate(row_values(ws, row_idx), start=1)
        if populated(value)
    }


def data_rows(ws, id_column: int = 1) -> list[int]:
    start = HEADER_ROWS[ws.title] + 1
    return [
        row_idx
        for row_idx in range(start, ws.max_row + 1)
        if populated(ws.cell(row=row_idx, column=id_column).value)
    ]


def has_print_area(ws) -> bool:
    return bool(str(ws.print_area or "").strip())


def validate(path: Path) -> dict[str, object]:
    errors: list[str] = []
    warnings: list[str] = []

    wb = load_openpyxl_workbook(path, data_only=False)

    if wb.sheetnames != EXPECTED_SHEETS:
        errors.append(f"Unexpected sheet order: {wb.sheetnames}")

    for sheet_name in EXPECTED_SHEETS:
        if sheet_name not in wb.sheetnames:
            errors.append(f"Missing sheet: {sheet_name}")

    if errors:
        return {"errors": errors, "warnings": warnings}

    for sheet_name, required_headers in REQUIRED_HEADERS.items():
        ws = wb[sheet_name]
        headers = header_map(ws)
        missing_headers = sorted(required_headers - set(headers))
        if missing_headers:
            errors.append(f"{sheet_name} missing headers: {', '.join(missing_headers)}")

    rm = wb["Requirement Matrix"]
    rm_headers = header_map(rm)
    requirement_rows = data_rows(rm)
    if not requirement_rows:
        errors.append("Requirement Matrix has no working requirement rows")
    for row_idx in requirement_rows:
        for header in CORE_REQUIREMENT_COLUMNS:
            col_idx = rm_headers.get(header)
            if col_idx and not populated(rm.cell(row=row_idx, column=col_idx).value):
                errors.append(f"Requirement Matrix row {row_idx} missing {header}")

    if not any(rm.cell(row=row_idx, column=rm_headers["Compliance Posture"]).value == "Gap" for row_idx in requirement_rows):
        warnings.append("Requirement Matrix has no gap rows; confirm this is intentional")

    ev = wb["Evaluation Crosswalk"]
    ev_headers = header_map(ev)
    eval_rows = data_rows(ev)
    if not eval_rows:
        errors.append("Evaluation Crosswalk has no factor rows")
    for row_idx in eval_rows:
        if not populated(ev.cell(row=row_idx, column=ev_headers["Related Req IDs"]).value):
            errors.append(f"Evaluation Crosswalk row {row_idx} missing Related Req IDs")

    dashboard = wb["Dashboard"]
    if not str(dashboard["A5"].value or "").startswith("=COUNTA"):
        errors.append("Dashboard total requirements formula is missing or changed")
    if not str(dashboard["D5"].value or "").startswith("=COUNTIFS"):
        errors.append("Dashboard mandatory rows formula is missing or changed")

    for sheet_name in EXPECTED_SHEETS:
        ws = wb[sheet_name]
        if not has_print_area(ws):
            errors.append(f"{sheet_name} missing print area")
        if ws.page_setup.fitToWidth != 1:
            errors.append(f"{sheet_name} is not configured to fit to one page wide")
        page_setup = ws.sheet_properties.pageSetUpPr
        if page_setup is None or page_setup.fitToPage is not True:
            errors.append(f"{sheet_name} is not configured for fit-to-page printing")

    return {
        "sheet_count": len(wb.sheetnames),
        "requirement_rows": len(requirement_rows),
        "evaluation_rows": len(eval_rows),
        "errors": errors,
        "warnings": warnings,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Validate a GovTribe proposal control workbook.")
    parser.add_argument("workbook", type=Path, help="Input .xlsx workbook")
    args = parser.parse_args()

    path = args.workbook
    if not path.exists():
        print(f"ERROR: workbook not found: {path}", file=sys.stderr)
        return 1

    summary = validate(path)
    print(json.dumps(summary, indent=2))

    return 1 if summary.get("errors") else 0


if __name__ == "__main__":
    raise SystemExit(main())
