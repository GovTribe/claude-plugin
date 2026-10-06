"""Layout-independent proposal controls from an explicit task/source contract."""

from datetime import datetime
from urllib.parse import urlparse
from delivery_evidence import load_contract, finish

FIELDS = {
    "tables",
    "requirement_ids",
    "deadline",
    "amendment_impacts",
    "source_classifications",
}
ROLES = {"requirements", "sources", "deadlines", "gaps", "amendments"}
COLUMNS = {
    "requirements": {
        "id",
        "source_id",
        "citation",
        "classification",
        "mandatory",
        "response",
        "owner",
    },
    "sources": {"id", "classification", "citation"},
    "deadlines": {"deadline"},
    "gaps": {"id", "status"},
    "amendments": {"id", "requirement_id"},
}


def instant(value):
    parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    if parsed.tzinfo is None:
        raise ValueError("Deadline requires an explicit timezone or UTC offset")
    return parsed


def read_contract(path):
    c = load_contract(path, FIELDS)
    if (
        not isinstance(c.get("tables"), list)
        or not c["tables"]
        or not isinstance(c.get("requirement_ids"), list)
        or not c["requirement_ids"]
    ):
        raise ValueError(
            "Proposal contract requires tables and nonempty source requirement_ids"
        )
    instant(c.get("deadline"))
    for table in c["tables"]:
        if (
            not isinstance(table, dict)
            or table.get("role") not in ROLES
            or not isinstance(table.get("sheet"), str)
            or not isinstance(table.get("header_row"), int)
            or table["header_row"] < 1
            or not isinstance(table.get("columns"), dict)
        ):
            raise ValueError("Invalid proposal table mapping")
        if not COLUMNS[table["role"]].issubset(table["columns"]):
            raise ValueError("Missing semantic columns for " + table["role"])
        if any(
            not isinstance(col, int) or col < 1 for col in table["columns"].values()
        ):
            raise ValueError("Columns require positive one-based indices")
    if not {"requirements", "sources", "deadlines"}.issubset(
        {t["role"] for t in c["tables"]}
    ):
        raise ValueError("Map requirements, sources, and deadlines")
    if len(set(c["requirement_ids"])) != len(c["requirement_ids"]):
        raise ValueError("Duplicate requirement IDs in source inventory")
    if not isinstance(c.get("amendment_impacts", {}), dict):
        raise ValueError(
            "amendment_impacts must map amendment IDs to requirement ID arrays"
        )
    for amendment, required in c.get("amendment_impacts", {}).items():
        if (
            not isinstance(amendment, str)
            or not amendment.strip()
            or not isinstance(required, list)
            or any(
                not isinstance(r, str) or r not in c["requirement_ids"]
                for r in required
            )
            or len(set(required)) != len(required)
        ):
            raise ValueError(
                "amendment_impacts requires amendment IDs and unique known requirement ID arrays"
            )
    return c


def rows(wb, t):
    ws = wb[t["sheet"]]
    return [
        {name: ws.cell(row, col).value for name, col in t["columns"].items()}
        for row in range(t["header_row"] + 1, ws.max_row + 1)
        if any(ws.cell(row, col).value is not None for col in t["columns"].values())
    ]


def portable(value):
    parsed = urlparse(str(value or ""))
    return parsed.scheme in ("https", "http") and bool(parsed.netloc)


def validate(path, c):
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "DEGRADED: openpyxl is unavailable. Contract parsing remains available. "
            "Use the host's spreadsheet capability or compare Markdown/CSV control "
            "tables against the independent source contract, preserving requirement "
            "and amendment inventories, citations, local deadlines and unresolved "
            "mandatory inputs. Workbook validation and rendering remain unverified."
        ) from exc
    wb = load_workbook(path, data_only=True)
    errors = []
    warnings = []
    unresolved = []
    data = {}
    for t in c["tables"]:
        data.setdefault(t["role"], []).extend(rows(wb, t))
    sources = {r["id"]: r for r in data["sources"]}
    if len(sources) != len(data["sources"]):
        errors.append("Duplicate source IDs")
    classifications = c.get(
        "source_classifications",
        [
            "solicitation",
            "amendment",
            "official-qa",
            "customer-input",
            "internal-control",
        ],
    )
    for row in data["sources"]:
        if row["classification"] not in classifications or not portable(
            row["citation"]
        ):
            errors.append(
                "Source needs classification and portable citation: " + str(row["id"])
            )
    requirements = data["requirements"]
    ids = [r["id"] for r in requirements]
    if len(set(ids)) != len(ids):
        errors.append("Duplicate requirement IDs")
    if set(ids) != set(c["requirement_ids"]):
        errors.append("Source requirement coverage mismatch")
    placeholders = ("tbd", "todo", "[insert", "[placeholder", "unknown", "missing")
    for row in requirements:
        label = str(row["id"])
        if row["source_id"] not in sources:
            errors.append("Unknown source ID for " + label)
        if row["classification"] not in ("source-requirement", "internal-control"):
            errors.append("Unclassified requirement/control: " + label)
        if not portable(row["citation"]):
            errors.append("Missing portable requirement citation: " + label)
        mandatory_flag = str(row["mandatory"]).strip().lower()
        if mandatory_flag not in (
            "yes",
            "true",
            "1",
            "mandatory",
            "no",
            "false",
            "0",
            "optional",
        ):
            errors.append("Unknown mandatory flag: " + label)
        mandatory = mandatory_flag in (
            "yes",
            "true",
            "1",
            "mandatory",
        )
        missing = not row["response"] or any(
            p in str(row["response"]).lower() for p in placeholders
        )
        if not row["owner"]:
            unresolved.append("missing-owner:" + label)
        if mandatory and missing:
            unresolved.append("mandatory-response:" + label)
    deadline = instant(c["deadline"])
    if not data["deadlines"]:
        errors.append("No delivered deadline controls")
    for row in data["deadlines"]:
        try:
            delivered = instant(row["deadline"])
            if delivered != deadline:
                errors.append("Conflicting active deadline: " + str(row["deadline"]))
            elif delivered.utcoffset() != deadline.utcoffset():
                errors.append(
                    "Deadline lost source-local time/UTC offset: " + str(row["deadline"])
                )
        except ValueError as exc:
            errors.append(str(exc))
    impacts = {}
    for row in data.get("amendments", []):
        impacted = impacts.setdefault(row["id"], set())
        if row["requirement_id"] not in (None, ""):
            impacted.add(row["requirement_id"])
    if set(impacts) != set(c.get("amendment_impacts", {})):
        errors.append(
            "Amendment inventory mismatch: delivered IDs must match amendment_impacts"
        )
    for amendment, required in c.get("amendment_impacts", {}).items():
        if impacts.get(amendment, set()) != set(required):
            errors.append("Amendment impact coverage mismatch: " + amendment)
    for row in data.get("gaps", []):
        if str(row["status"]).lower() not in ("resolved", "closed"):
            unresolved.append("open-gap:" + str(row["id"]))
    state = c.get("state", "submission-candidate")
    if unresolved and state == "submission-candidate":
        errors.append(
            "Submission readiness failed: unresolved mandatory inputs or controls"
        )
    elif unresolved:
        warnings.append("Draft has explicit unresolved inputs; not submission-ready")
    result = {
        "errors": errors,
        "warnings": warnings,
        "sheet_count": len(wb.sheetnames),
        "requirement_rows": len(requirements),
        "evaluation_rows": 0,
    }
    return finish(
        result,
        path,
        c,
        [
            "requirement-source-coverage",
            "source-classification-citations",
            "timezone-deadlines",
            "source-local-deadlines",
            "amendment-impacts",
            "mandatory-readiness",
        ],
        unresolved,
    )
