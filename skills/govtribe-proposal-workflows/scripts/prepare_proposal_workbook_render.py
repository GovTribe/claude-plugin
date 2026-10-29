#!/usr/bin/env python3
"""Prepare a GovTribe proposal control workbook for PDF rendering.

The bundled workbook is optimized for editing. Before visual QA or PDF delivery,
set print areas to populated ranges, fit wide sheets to one page, and expand
rows that contain long wrapped text so LibreOffice does not clip content.

For selected revisions, --previous-workbook grows changed wrapped rows within
the retained print panels without replacing layout or resaving workbook objects.
"""
from __future__ import annotations

import argparse
import math
import os
import posixpath
import re
import shutil
import sys
import tempfile
import textwrap
import unicodedata
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

SHEET_COLUMNS = {
    "Start Here": 7,
    "Opportunity Setup": 9,
    "Dashboard": 11,
    "Requirement Matrix": 21,
    "Evaluation Crosswalk": 12,
    "Submission Checklist": 11,
    "Pricing & Deliverables": 10,
    "Questions-Risks": 11,
    "Amendment Log": 10,
    "Sources": 5,
    "Lists": 18,
}

load_workbook = None
get_column_letter = None
PageMargins = None
PageSetupProperties = None


def load_openpyxl() -> int:
    global load_workbook, get_column_letter, PageMargins, PageSetupProperties

    try:
        from openpyxl import load_workbook as openpyxl_load_workbook
        from openpyxl.utils import get_column_letter as openpyxl_get_column_letter
        from openpyxl.worksheet.page import PageMargins as OpenPyxlPageMargins
        from openpyxl.worksheet.properties import (
            PageSetupProperties as OpenPyxlPageSetupProperties,
        )
    except ModuleNotFoundError:
        print(
            "DEGRADED: openpyxl is unavailable. Use the external host's spreadsheet "
            "rendering capability and manually verify print areas, fit-to-width, and "
            "wrapped row heights. Deliver Markdown or CSV control views if rendering "
            "is unavailable.",
            file=sys.stderr,
        )
        return 2

    load_workbook = openpyxl_load_workbook
    get_column_letter = openpyxl_get_column_letter
    PageMargins = OpenPyxlPageMargins
    PageSetupProperties = OpenPyxlPageSetupProperties

    return 0


MIN_PRINT_ROWS = {
    "Start Here": 31,
    "Opportunity Setup": 14,
    "Dashboard": 34,
    "Requirement Matrix": 8,
    "Evaluation Crosswalk": 8,
    "Submission Checklist": 7,
    "Pricing & Deliverables": 7,
    "Questions-Risks": 7,
    "Amendment Log": 7,
    "Sources": 8,
    "Lists": 16,
}

LONG_TEXT_SHEETS = {
    "Requirement Matrix",
    "Evaluation Crosswalk",
    "Submission Checklist",
    "Pricing & Deliverables",
    "Questions-Risks",
    "Amendment Log",
    "Sources",
}

ROW_ID_PRINT_SHEETS = {
    "Requirement Matrix",
    "Evaluation Crosswalk",
    "Submission Checklist",
    "Pricing & Deliverables",
    "Questions-Risks",
    "Amendment Log",
}


def is_populated(value: object) -> bool:
    return value is not None and str(value).strip() != ""


def last_populated_row(ws, max_col: int) -> int:
    last_row = 1
    if ws.title in ROW_ID_PRINT_SHEETS:
        for row_idx in range(1, ws.max_row + 1):
            if is_populated(ws.cell(row=row_idx, column=1).value):
                last_row = row_idx

        return max(last_row, MIN_PRINT_ROWS.get(ws.title, last_row))

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=max_col):
        if any(is_populated(cell.value) for cell in row):
            last_row = row[0].row

    return max(last_row, MIN_PRINT_ROWS.get(ws.title, last_row))


def effective_column_width(ws, column: int) -> float:
    width = ws.column_dimensions[get_column_letter(column)].width

    return float(width or 10)


def estimated_lines(value: object, width: float) -> int:
    if not is_populated(value):
        return 1

    text = str(value)
    explicit_lines = text.count("\n") + 1
    wrapped_lines = math.ceil(len(text) / max(width * 1.1, 8))

    return max(explicit_lines, wrapped_lines, 1)


def adjust_row_heights(ws, max_col: int, max_row: int) -> None:
    if ws.title not in LONG_TEXT_SHEETS:
        return

    for row_idx in range(1, max_row + 1):
        max_lines = 1
        for col_idx in range(1, max_col + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            max_lines = max(
                max_lines,
                estimated_lines(cell.value, effective_column_width(ws, col_idx)),
            )

        if max_lines > 1:
            ws.row_dimensions[row_idx].height = min(max(18, max_lines * 15), 120)


def revision_column_width(ws, column: int) -> float:
    # Grouped column definitions need not have an entry for every column letter.
    for dimension in ws.column_dimensions.values():
        if dimension.min <= column <= dimension.max:
            return 0 if dimension.hidden else float(dimension.width)
    return float(ws.sheet_format.defaultColWidth or 8.43)


def revision_text_height(ws, cell) -> float:
    width = revision_column_width(ws, cell.column)
    if width == 0:
        return 0
    for merged in ws.merged_cells.ranges:
        if cell.coordinate in merged:
            if merged.min_row != merged.max_row:
                raise ValueError(f"{ws.title}!{cell.coordinate}: vertically merged text needs native row sizing")
            width = sum(
                revision_column_width(ws, col)
                for col in range(merged.min_col, merged.max_col + 1)
            )
            break
    if cell.alignment.textRotation:
        raise ValueError(f"{ws.title}!{cell.coordinate}: rotated text needs native row sizing")

    size = float(cell.font.sz or 11)
    # Excel column units approximate 7px at the normal 11pt font. Leave padding,
    # scale for the cell font, and wrap at words instead of dividing total length.
    font_scale = size / 11 * (1.1 if cell.font.bold else 1)
    indent = int(cell.alignment.indent or 0) * 3
    capacity = max(1, math.floor((width * 7 - 10) / (7 * font_scale)) - indent)
    lines = 0
    for paragraph in cell.value.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        text = "".join(
            "" if unicodedata.combining(char)
            else "xx" if unicodedata.east_asian_width(char) in ("W", "F")
            else char
            for char in paragraph
        )
        lines += max(1, len(textwrap.wrap(text, width=capacity)))
    height = math.ceil(lines * max(12, size * 1.2) + 3)
    if height > 409.5:
        raise ValueError(f"{ws.title}!{cell.coordinate}: text exceeds Excel's row-height limit; review the layout natively")
    return height


def revision_row_heights(workbook, previous) -> dict[str, dict[int, float]]:
    from openpyxl.worksheet.print_settings import PrintArea

    if workbook.sheetnames != previous.sheetnames:
        raise ValueError("Selected previous workbook must have the same worksheet names and order")
    changes = {}
    for ws in workbook:
        if (
            ws.title not in LONG_TEXT_SHEETS
            or ws.sheet_state != "visible"
            or not ws.print_area
        ):
            continue
        areas = PrintArea.from_string(ws.print_area).ranges
        for row in ws:
            for cell in row:
                if (
                    not isinstance(cell.value, str)
                    or not cell.value.strip()
                    or cell.data_type == "f"
                    or not cell.alignment.wrap_text
                    or cell.value == previous[ws.title][cell.coordinate].value
                    or ws.row_dimensions[cell.row].hidden
                    or not any(cell.coordinate in area for area in areas)
                ):
                    continue
                height = revision_text_height(ws, cell)
                current = (
                    ws.row_dimensions[cell.row].height
                    or ws.sheet_format.defaultRowHeight or 15
                )
                if height > current:
                    rows = changes.setdefault(ws.title, {})
                    rows[cell.row] = max(rows.get(cell.row, current), height)
    return changes


def patch_row_heights(data: bytes, heights: dict[int, float]) -> bytes:
    # Keep all cell XML (including cached values) and non-height row attributes
    # byte-identical. Parsing/reserializing the worksheet would rewrite them.
    row_tag = re.compile(rb'<(?:[A-Za-z_][\w.-]*:)?row(?=[\s/>])(?:[^\'">]|\'[^\']*\'|"[^"]*")*>')
    found = set()

    def attribute(tag: bytes, name: bytes, value: bytes) -> bytes:
        pattern = rb'(?<![\w:.-])' + name + rb'\s*=\s*([\'"])[^\'"]*\1'
        replacement = name + b'="' + value + b'"'
        if re.search(pattern, tag):
            return re.sub(pattern, lambda _: replacement, tag, count=1)
        end = 2 if tag.endswith(b"/>") else 1
        return tag[:-end] + b" " + replacement + tag[-end:]

    def replace(match):
        tag = match.group()
        index = re.search(rb'(?<![\w:.-])r\s*=\s*([\'"])(\d+)\1', tag)
        if index is None or int(index[2]) not in heights:
            return tag
        number = int(index[2])
        if number in found:
            raise ValueError(f"Duplicate worksheet row {number}")
        found.add(number)
        tag = attribute(tag, b"ht", f"{heights[number]:g}".encode())
        return attribute(tag, b"customHeight", b"1")

    patched = row_tag.sub(replace, data)
    if found != set(heights):
        raise ValueError("Cannot safely locate every revised worksheet row")
    return patched


def save_revision_rows(source: Path, output: Path, changes: dict[str, dict[int, float]]) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    if not changes:
        if source.resolve() != output.resolve():
            shutil.copyfile(source, output)
        return
    namespace = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with zipfile.ZipFile(source) as archive:
        if len(archive.namelist()) != len(set(archive.namelist())):
            raise ValueError("Duplicate workbook package members are unsupported")
        relationships = {
            node.get("Id"): node.get("Target")
            for node in ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            if node.get("TargetMode") != "External"
        }
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        parts = {}
        for sheet in workbook.findall("s:sheets/s:sheet", namespace):
            if sheet.get("name") not in changes:
                continue
            relationship_id = sheet.get(
                "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
            )
            target = relationships[relationship_id]
            part = posixpath.normpath(
                target.lstrip("/") if target.startswith("/")
                else posixpath.join("xl", target)
            )
            if not part.startswith("xl/") or part in parts:
                raise ValueError("Unsupported worksheet package relationship")
            parts[part] = changes[sheet.get("name")]
        if len(parts) != len(changes):
            raise ValueError("Cannot locate every revised worksheet part")
        with tempfile.NamedTemporaryFile(
            dir=output.parent, suffix=".xlsx", delete=False
        ) as temporary:
            temporary_path = Path(temporary.name)
        try:
            with zipfile.ZipFile(temporary_path, "w") as revised:
                revised.comment = archive.comment
                for member in archive.infolist():
                    data = archive.read(member.filename)
                    if member.filename in parts:
                        data = patch_row_heights(data, parts[member.filename])
                    revised.writestr(member, data)
            os.replace(temporary_path, output)
        finally:
            temporary_path.unlink(missing_ok=True)


def prepare_sheet(ws) -> None:
    max_col = SHEET_COLUMNS.get(ws.title, ws.max_column)
    for row in ws.iter_rows():
        for cell in row:
            if is_populated(cell.value):
                max_col = max(max_col, cell.column)
    for merged in ws.merged_cells.ranges:
        if is_populated(ws.cell(merged.min_row, merged.min_col).value):
            max_col = max(max_col, merged.max_col)
    max_row = last_populated_row(ws, max_col)
    last_col = get_column_letter(max_col)

    ws.print_area = f"A1:{last_col}{max_row}"
    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.scale = None
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins = PageMargins(
        left=0.25, right=0.25, top=0.4, bottom=0.4, header=0.2, footer=0.2
    )

    adjust_row_heights(ws, max_col, max_row)


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Prepare a GovTribe proposal workbook for visual rendering."
    )
    parser.add_argument("workbook", type=Path, help="Input .xlsx workbook")
    parser.add_argument(
        "--output",
        type=Path,
        help="Output .xlsx path. Defaults to overwriting with --in-place.",
    )
    parser.add_argument(
        "--in-place", action="store_true", help="Overwrite the input workbook"
    )
    parser.add_argument(
        "--profile", choices=("eleven-sheet",), default="eleven-sheet"
    )
    parser.add_argument(
        "--previous-workbook", type=Path,
        help="Selected original: grow only changed wrapped rows in retained print panels; preserve package contents and print settings",
    )
    args = parser.parse_args()
    if not args.workbook.exists():
        print(f"ERROR: workbook not found: {args.workbook}")
        return 1

    if not args.in_place and args.output is None:
        print("ERROR: provide --output or --in-place")
        return 2

    if args.in_place and args.output is not None:
        parser.error("Use either --output or --in-place, not both")

    output = args.workbook if args.in_place else args.output
    if args.previous_workbook is not None:
        if not args.previous_workbook.is_file():
            parser.error("Selected previous workbook does not exist")
        if any(
            path.resolve() == args.previous_workbook.resolve()
            or (path.exists() and path.samefile(args.previous_workbook))
            for path in (args.workbook, output)
        ):
            parser.error("Keep the selected previous workbook separate from the edited input and output")

    if exit_code := load_openpyxl():
        return exit_code

    wb = load_workbook(args.workbook)

    if set(wb.sheetnames) != set(SHEET_COLUMNS):
        print(
            "ERROR: This helper supports only the eleven-sheet template. Use native spreadsheet tools for custom layouts.",
            file=sys.stderr,
        )
        return 1

    if args.previous_workbook is not None:
        try:
            changes = revision_row_heights(wb, load_workbook(args.previous_workbook))
            save_revision_rows(args.workbook, output, changes)
        except (ValueError, KeyError, OSError, ET.ParseError, zipfile.BadZipFile) as exc:
            print(f"ERROR: {exc}", file=sys.stderr)
            return 1
        print(f"Saved revision with retained print settings: {output}")
        summary = ", ".join(
            f"{sheet}!{row}={height:g}pt"
            for sheet, rows in changes.items()
            for row, height in sorted(rows.items())
        )
        print("Rows enlarged: " + (summary or "none"))
        print("Row sizes are estimates; inspect the saved workbook's original print layout before delivery.")
        return 0

    prepared = []
    for ws in wb.worksheets:
        prepare_sheet(ws)
        prepared.append(ws.title)

    wb.save(output)
    print(f"Prepared workbook render layout: {output}")
    print(f"Sheets prepared: {', '.join(prepared)}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
